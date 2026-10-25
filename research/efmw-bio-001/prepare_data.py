"""Extract the authors' unmodified spreadsheets and compute protocol-fixed features."""
from pathlib import Path
import json, hashlib, datetime, zipfile
import numpy as np
import openpyxl
from lxml import etree

ROOT=Path(__file__).resolve().parent
def numeric_xlsx(path):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    sheet=wb.active
    nr,nc=sheet.max_row-1,sheet.max_column
    header=list(next(sheet.values));wb.close()
    result=np.full((nr,nc),np.nan,dtype=np.float64)
    ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    with zipfile.ZipFile(path) as z, z.open('xl/worksheets/sheet1.xml') as f:
        for _,row in etree.iterparse(f,events=('end',),tag=ns+'row'):
            ri=int(row.get('r'))-2
            if ri>=0:
                for cell in row:
                    v=cell.find(ns+'v')
                    if v is None or v.text is None: continue
                    ref=cell.get('r');ci=0
                    for ch in ref:
                        if ch.isalpha():ci=ci*26+ord(ch)-64
                        else:break
                    if cell.get('t') not in ['s','str','inlineStr','e']:
                        result[ri,ci-1]=float(v.text)
            row.clear()
            while row.getprevious() is not None:del row.getparent()[0]
    return result,header

def memory(x,a):
    state=np.zeros(len(x))
    for t in range(1,x.shape[1]): state=a*state+(1-a)*x[:,t]
    return state

def main():
    freeze=json.loads((ROOT/'protocol_freeze.json').read_text())
    assert hashlib.sha256((ROOT/'PROTOCOL.md').read_bytes()).hexdigest()==freeze['protocol_sha256']
    print('Reading ERK traces',flush=True)
    X,xcols=numeric_xlsx(ROOT/'data/df_erk_MCF10A.xlsx')
    print('Reading protein outcomes after protocol freeze',flush=True)
    Y,ycols=numeric_xlsx(ROOT/'data/df_ETGs_MCF10A.xlsx')
    wb=openpyxl.load_workbook(ROOT/'data/df_tx_MCF10A.xlsx',read_only=True,data_only=True)
    rows=list(wb.active.values);wb.close()
    meta=rows[1:]
    assert len(meta)==len(X)==len(Y) and X.shape[1]==190
    group=np.array([r[2] for r in meta],dtype=int)
    well=np.array([r[1] for r in meta],dtype=int)
    condition=np.array([r[0] for r in meta],dtype=str)
    good=np.isfinite(X).sum(axis=1)>=152
    row_ids=np.arange(len(X))+2
    stats={'original_cells':len(X),'time_samples':X.shape[1],
           'trace_qc_pass':int(good.sum()),'trace_qc_excluded':int((~good).sum()),
           'timepoints_missing_before_interpolation':int((~np.isfinite(X[good])).sum()),
           'replicate_counts_before_qc':{str(g):int((group==g).sum()) for g in np.unique(group)},
           'parsed_at_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),
           'target_columns':ycols,'source_row_count_differs_from_paper_figure_count':True,
           'source_row_count_note':'The downloadable tables have 71,459 rows. Paper Figure 2 reports 97,960 cells; reconciliation of that difference is not established.'}
    X,Y,group,well,condition,row_ids=X[good],Y[good],group[good],well[good],condition[good],row_ids[good]
    ts=np.arange(X.shape[1])
    for i in np.flatnonzero(~np.isfinite(X).all(axis=1)):
        valid=np.isfinite(X[i]);X[i]=np.interp(ts,ts[valid],X[i,valid])
    base=np.column_stack([X.mean(axis=1),X.std(axis=1),X[:,:10].mean(axis=1),X[:,-10:].mean(axis=1)])
    efmw=np.abs(memory(X,.97))
    powers=(.97**np.arange(X.shape[1]-2,-1,-1))*.03
    ewma=np.abs(X[:,1:]@powers)
    np.testing.assert_allclose(efmw,ewma,rtol=1e-12,atol=1e-12)
    stats['memory_equivalence_max_abs_error']=float(np.max(np.abs(efmw-ewma)))
    cascade=[]
    for a in [.8,.97,.995]:
        for b in [.8,.97,.995]:
            m=np.zeros(len(X));p=np.zeros(len(X))
            for t in range(1,X.shape[1]):
                m=a*m+(1-a)*X[:,t];p=b*p+(1-b)*m
            cascade.append(p)
    models={
        'intercept':np.empty((len(X),0)),
        'snapshot':base[:,3:4],
        'time_average':base[:,0:1],
        'summary_ablation':base,
        'efmw_memory':np.column_stack([base,efmw]),
        'matched_ewma':np.column_stack([base,ewma]),
        'multiscale_ema':np.column_stack([base,*[memory(X,a) for a in [.5,.8,.9,.97,.99,.995]]]),
        'temporal_bins':X.reshape(len(X),19,10).mean(axis=2),
        'production_decay':np.column_stack([base,*cascade])}
    stats['target_usable_counts']={c:int((np.isfinite(Y[:,j])&(Y[:,j]>0)).sum()) for j,c in enumerate(ycols)}
    stats['predictor_counts']={k:v.shape[1] for k,v in models.items()}
    np.savez_compressed(ROOT/'data/analysis_data.npz',X=X,Y=Y,replicate=group,well=well,condition=condition,
                        source_excel_row=row_ids,target_names=np.array(ycols),**{'features_'+k:v for k,v in models.items()})
    (ROOT/'results/data_audit.json').write_text(json.dumps(stats,indent=2)+'\n')
    print(json.dumps(stats,indent=2),flush=True)

if __name__=='__main__':main()
